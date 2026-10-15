"""Build and read the spreadsheet and Word menu templates."""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

Document = Workbook = load_workbook = Font = PatternFill = None
DataValidation = WD_ORIENT = None
try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # Keep the rest of the app available until extra packages are installed.
    pass


MENU_HEADERS = (
    "Nombre",
    "Precio (S/)",
    "Tipo de menú",
    "Fecha (AAAA-MM-DD)",
    "Categoría",
    "Horario",
    "Descripción actual",
    "Ingredientes",
    "Información dietética",
    "Tiempo de preparación (min)",
    "Nivel de picante (0-5)",
)
WORD_MENU_HEADERS = MENU_HEADERS[:9]
MAX_MENU_ROWS = 40
MAX_MENU_TEXT = 30000
MAX_ARCHIVE_MEMBERS = 500
MAX_ARCHIVE_UNPACKED_BYTES = 25 * 1024 * 1024


class MenuDocumentError(ValueError):
    """The uploaded document is empty, malformed, or outside supported limits."""


def _require_document_packages() -> None:
    if any(package is None for package in (Document, Workbook, load_workbook, Font, PatternFill, DataValidation, WD_ORIENT)):
        raise MenuDocumentError("Instala las dependencias openpyxl y python-docx para usar plantillas Excel y Word.")


def build_menu_template(file_format: str) -> BytesIO:
    """Return a blank XLSX or DOCX template as a rewound in-memory stream."""
    _require_document_packages()
    if file_format == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Carta"
        sheet.append(MENU_HEADERS)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="FF385C")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:K{MAX_MENU_ROWS + 1}"
        menu_type_validation = DataValidation(type="list", formula1='"carta,variable"', allow_blank=False)
        category_validation = DataValidation(type="list", formula1='"Entrada,Fondo,Postre,Bebida"', allow_blank=True)
        meal_validation = DataValidation(type="list", formula1='"desayuno,almuerzo,cena"', allow_blank=True)
        menu_date_validation = DataValidation(
            type="custom",
            formula1='=AND($C2<>"",OR($C2="carta",AND($C2="variable",OR(ISNUMBER($D2),AND(ISTEXT($D2),LEN($D2)=10,MID($D2,5,1)="-",MID($D2,8,1)="-")))))',
            allow_blank=False,
        )
        menu_type_validation.error = "Selecciona carta o variable."
        menu_type_validation.errorTitle = "Tipo de menú requerido"
        menu_type_validation.showInputMessage = True
        menu_type_validation.promptTitle = "Tipo de menú"
        menu_type_validation.prompt = "Elige carta o variable."
        menu_date_validation.error = "Para menú variable, ingresa una fecha válida."
        menu_date_validation.errorTitle = "Fecha requerida"
        menu_date_validation.showInputMessage = True
        menu_date_validation.promptTitle = "Fecha del menú variable"
        menu_date_validation.prompt = "Usa el formato AAAA-MM-DD, por ejemplo 2026-10-15."
        for validation in (menu_type_validation, category_validation, meal_validation, menu_date_validation):
            sheet.add_data_validation(validation)
        menu_type_validation.add(f"C2:C{MAX_MENU_ROWS + 1}")
        menu_date_validation.add(f"D2:D{MAX_MENU_ROWS + 1}")
        category_validation.add(f"E2:E{MAX_MENU_ROWS + 1}")
        meal_validation.add(f"F2:F{MAX_MENU_ROWS + 1}")
        for row in range(2, MAX_MENU_ROWS + 2):
            sheet[f"D{row}"].number_format = "yyyy-mm-dd"
        for column, width in zip("ABCDEFGHIJK", (28, 16, 24, 24, 18, 18, 48, 36, 32, 25, 22)):
            sheet.column_dimensions[column].width = width

        instructions = workbook.create_sheet("Instrucciones")
        instructions.append(("Cómo completar la plantilla",))
        instructions.append(("Agrega un plato por fila en la hoja Carta.",))
        instructions.append(("Nombre y precio son obligatorios. El precio debe ser numérico.",))
        instructions.append(("Tipo de menú es obligatorio: carta o variable.",))
        instructions.append(("Si eliges variable, completa la fecha del menú en formato AAAA-MM-DD.",))
        instructions.append(("Categoría: Entrada, Fondo, Postre o Bebida.",))
        instructions.append(("Horario: desayuno, almuerzo o cena. Se puede dejar vacío.",))
        instructions.append(("Separa ingredientes e información dietética con comas.",))
        instructions.append(("Las descripciones sugeridas por Qwen se revisan antes de guardar.",))

        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)
        return stream

    if file_format == "docx":
        document = Document()
        section = document.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
        document.add_heading("Plantilla de carta del restaurante", 0)
        document.add_paragraph(
            "Completa una fila por plato. Nombre y precio son obligatorios. "
            "Indica si es de carta o variable; para un plato variable, completa la fecha del menú (AAAA-MM-DD). "
            "Agrega más filas si las necesitas. Deja en blanco los datos que no tengas; podrás revisar las sugerencias de Qwen antes de guardarlas."
        )
        word_headers = WORD_MENU_HEADERS
        table = document.add_table(rows=9, cols=len(word_headers))
        table.style = "Table Grid"
        for cell, header in zip(table.rows[0].cells, word_headers):
            cell.text = header
        document.add_paragraph("Categorías: Entrada, Fondo, Postre o Bebida.")
        document.add_paragraph("Horarios: desayuno, almuerzo o cena.")
        document.add_paragraph("Tipo de menú: carta o variable. Para variable, la fecha es obligatoria y debe usar AAAA-MM-DD.")
        document.add_paragraph("Separa ingredientes e información dietética con comas.")
        stream = BytesIO()
        document.save(stream)
        stream.seek(0)
        return stream

    raise MenuDocumentError("Formato de plantilla no compatible.")


def _header_key(value: object) -> str:
    text = " ".join(str(value or "").strip().lower().replace("\n", " ").split())
    return text.replace("á", "a").replace("é", "e").replace("í", "i").replace("ó", "o").replace("ú", "u")


_HEADER_ALIASES = {
    "name": {"nombre", "nombre del plato", "plato", "dish", "dish name"},
    "price": {"precio", "precio (s/)", "precio s/", "price", "price (s/)"},
    "menu_type": {"tipo de menu", "tipo de menu (carta/variable)", "carta o variable", "menu type"},
    "menu_date": {"fecha", "fecha (aaaa-mm-dd)", "fecha del menu", "fecha del menu (aaaa-mm-dd)", "menu date", "date"},
    "category": {"categoria", "tipo", "tipo de plato", "category"},
    "meal_time": {"horario", "tiempo de comida", "meal time"},
    "description": {"descripcion actual", "descripcion", "description"},
    "ingredients": {"ingredientes", "ingredients"},
    "dietary_info": {"informacion dietetica", "info dietetica", "dietary info"},
    "preparation_time": {"tiempo de preparacion (min)", "tiempo de preparacion", "preparation time"},
    "spiciness": {"nivel de picante (0-5)", "nivel de picante", "picante", "spiciness"},
}
_HEADER_LOOKUP = {alias: field for field, aliases in _HEADER_ALIASES.items() for alias in aliases}


def _canonical_header(value: object) -> str | None:
    return _HEADER_LOOKUP.get(_header_key(value))


def _clean_cell(value: object, limit: int = 1000) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())[:limit]


def _normalize_menu_type(value: object) -> str | None:
    normalized = _header_key(value)
    aliases = {
        "carta": "carta",
        "a la carta": "carta",
        "carta permanente": "carta",
        "permanente": "carta",
        "variable": "variable",
        "menu variable": "variable",
        "diario": "variable",
    }
    return aliases.get(normalized)


def _normalize_menu_date(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _clean_cell(value, limit=40)
    if not text:
        return None
    for date_format in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, date_format).date().isoformat()
        except ValueError:
            continue
    return None


def _rows_to_records(rows) -> list[dict]:
    rows = list(rows)
    if not rows:
        raise MenuDocumentError("La plantilla no contiene una tabla con platos.")

    header_row_index = None
    fields = []
    for index, row in enumerate(rows[:5]):
        candidate_fields = [_canonical_header(value) for value in row]
        if "name" in candidate_fields and "price" in candidate_fields:
            header_row_index = index
            fields = candidate_fields
            break
    if header_row_index is None:
        raise MenuDocumentError("No encontré las columnas Nombre y Precio de la plantilla.")
    has_menu_type_column = "menu_type" in fields

    records = []
    total_chars = 0
    for row_number, row in enumerate(rows[header_row_index + 1:], start=header_row_index + 2):
        values = {}
        for field, value in zip(fields, row):
            if field:
                values[field] = value if field == "menu_date" else _clean_cell(value)
        if not any(values.values()):
            continue
        if not values.get("name") and not values.get("price"):
            continue
        menu_type = _normalize_menu_type(values.get("menu_type")) if has_menu_type_column else "carta"
        if has_menu_type_column and not menu_type:
            raise MenuDocumentError(f"Indica si {values.get('name') or 'el plato'} es de carta o variable (fila {row_number}).")
        if not menu_type:
            raise MenuDocumentError(f"El tipo de menú de {values.get('name') or 'el plato'} debe ser carta o variable (fila {row_number}).")
        menu_date = _normalize_menu_date(values.get("menu_date"))
        if menu_type == "variable" and not menu_date:
            raise MenuDocumentError(f"Indica una fecha válida para el plato variable {values.get('name') or ''} (fila {row_number}).")

        records.append({
            "name": values.get("name", "")[:120],
            "price": values.get("price", "")[:40],
            "menu_type": menu_type,
            "menu_date": menu_date if menu_type == "variable" else None,
            "category": values.get("category", "")[:60],
            "meal_time": values.get("meal_time", "")[:30],
            "description": values.get("description", "")[:500],
            "ingredients": values.get("ingredients", "")[:500],
            "dietary_info": values.get("dietary_info", "")[:300],
            "preparation_time": values.get("preparation_time", "")[:20],
            "spiciness": values.get("spiciness", "")[:10],
        })
        total_chars += sum(len(str(value or "")) for value in records[-1].values())
        if len(records) > MAX_MENU_ROWS:
            raise MenuDocumentError(f"La plantilla admite hasta {MAX_MENU_ROWS} platos por carga.")
        if total_chars > MAX_MENU_TEXT:
            raise MenuDocumentError("El texto de la plantilla supera el límite permitido.")

    if not records:
        raise MenuDocumentError("Agrega al menos un plato debajo de los encabezados.")
    for index, record in enumerate(records, start=1):
        if not record["name"]:
            raise MenuDocumentError(f"Falta el nombre del plato en la fila {index + 1}.")
        if not record["price"]:
            raise MenuDocumentError(f"Falta el precio de {record['name']}.")
        try:
            price_text = record["price"].lower().replace("s/.", "").replace("s/", "").strip()
            price = float(price_text.replace(",", "."))
        except ValueError as exc:
            raise MenuDocumentError(f"El precio de {record['name']} debe ser numérico.") from exc
        if not 0 <= price <= 100000:
            raise MenuDocumentError(f"El precio de {record['name']} debe estar entre 0 y 100000.")
        record["price"] = round(price, 2)
    return records


def read_menu_document(file_stream, filename: str) -> list[dict]:
    """Extract bounded, normalized menu rows from XLSX or DOCX in memory."""
    _require_document_packages()
    suffix = Path(filename or "").suffix.lower()
    file_stream.seek(0)
    try:
        with ZipFile(file_stream) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise MenuDocumentError("El archivo contiene demasiados elementos para procesarlo.")
            if sum(member.file_size for member in members) > MAX_ARCHIVE_UNPACKED_BYTES:
                raise MenuDocumentError("El contenido descomprimido supera el límite permitido.")
    except MenuDocumentError:
        raise
    except (BadZipFile, OSError) as exc:
        raise MenuDocumentError("El archivo no es un Excel o Word válido. Descarga una plantilla nueva.") from exc
    finally:
        file_stream.seek(0)

    if suffix == ".xlsx":
        workbook = None
        try:
            workbook = load_workbook(file_stream, read_only=True, data_only=True)
            sheet = workbook["Carta"] if "Carta" in workbook.sheetnames else workbook.active
            rows = sheet.iter_rows(values_only=True)
            return _rows_to_records(rows)
        except MenuDocumentError:
            raise
        except Exception as exc:
            raise MenuDocumentError("No pude leer el archivo Excel. Descarga una plantilla nueva e inténtalo otra vez.") from exc
        finally:
            if workbook is not None:
                workbook.close()
    if suffix == ".docx":
        try:
            document = Document(file_stream)
            if not document.tables:
                raise MenuDocumentError("El documento Word no contiene la tabla de la plantilla.")
            rows = ([cell.text for cell in row.cells] for row in document.tables[0].rows)
            return _rows_to_records(rows)
        except MenuDocumentError:
            raise
        except Exception as exc:
            raise MenuDocumentError("No pude leer el archivo Word. Descarga una plantilla nueva e inténtalo otra vez.") from exc
    raise MenuDocumentError("Sube un archivo .xlsx o .docx.")

